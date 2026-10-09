import os
import re
import json
import time
import logging
import hashlib
from urllib.parse import urljoin, urlparse, parse_qs

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================

BASE_URL = "https://venuehunt.pk"
VENUES_URL = f"{BASE_URL}/venues"

OUTPUT_DIR = "venuehunt_output"
IMAGE_DIR = os.path.join(OUTPUT_DIR, "images")

EXCEL_FILE = os.path.join(
    OUTPUT_DIR,
    "venuehunt_venues.xlsx"
)

PROGRESS_FILE = os.path.join(
    OUTPUT_DIR,
    "progress.json"
)

ERROR_LOG = os.path.join(
    OUTPUT_DIR,
    "errors.log"
)

REQUEST_DELAY = 1.0
MAX_RETRIES = 3
REQUEST_TIMEOUT = 30

# Set to None to scrape all available pages.
MAX_PAGES = None

# Set True if you want images downloaded.
DOWNLOAD_IMAGES = True


# ============================================================
# DIRECTORY SETUP
# ============================================================

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(IMAGE_DIR, exist_ok=True)


# ============================================================
# LOGGING
# ============================================================

logging.basicConfig(
    filename=ERROR_LOG,
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s"
)

console = logging.StreamHandler()
console.setLevel(logging.INFO)

logging.getLogger().addHandler(console)


# ============================================================
# HTTP SESSION
# ============================================================

session = requests.Session()

session.headers.update({
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/151.0.0.0 Safari/537.36"
    ),
    "Accept": (
        "text/html,application/xhtml+xml,"
        "application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8"
    ),
    "Accept-Language": "en-US,en;q=0.9",
    "Connection": "keep-alive",
})


# ============================================================
# HELPERS
# ============================================================

def clean_text(value):
    if value is None:
        return ""

    value = value.replace("\xa0", " ")

    return re.sub(
        r"\s+",
        " ",
        value
    ).strip()


def normalize_url(url):
    if not url:
        return ""

    url = urljoin(BASE_URL, url)

    # Remove query string and fragment
    parsed = urlparse(url)

    clean = (
        f"{parsed.scheme}://"
        f"{parsed.netloc}"
        f"{parsed.path}"
    )

    if not clean.endswith("/"):
        clean += "/"

    return clean


def safe_filename(value, max_length=100):

    value = clean_text(value)

    value = re.sub(
        r"[^\w\s-]",
        "",
        value
    )

    value = re.sub(
        r"\s+",
        "_",
        value
    )

    return value[:max_length] or "venue"


def unique_list(items):

    result = []
    seen = set()

    for item in items:

        item = clean_text(item)

        if not item:
            continue

        key = item.lower()

        if key not in seen:
            seen.add(key)
            result.append(item)

    return result


def request_get(url, stream=False):

    last_error = None

    for attempt in range(
        1,
        MAX_RETRIES + 1
    ):

        try:

            response = session.get(
                url,
                timeout=REQUEST_TIMEOUT,
                stream=stream
            )

            response.raise_for_status()

            return response

        except Exception as exc:

            last_error = exc

            logging.warning(
                "Request failed [%s/%s]: %s | %s",
                attempt,
                MAX_RETRIES,
                url,
                exc
            )

            if attempt < MAX_RETRIES:

                time.sleep(
                    2 ** attempt
                )

    raise last_error


def get_soup(url):

    response = request_get(url)

    return BeautifulSoup(
        response.text,
        "lxml"
    )


# ============================================================
# PROGRESS / RESUME
# ============================================================

def load_progress():

    if not os.path.exists(PROGRESS_FILE):
        return {
            "scraped_urls": [],
            "venues": []
        }

    try:

        with open(
            PROGRESS_FILE,
            "r",
            encoding="utf-8"
        ) as f:

            return json.load(f)

    except Exception:

        return {
            "scraped_urls": [],
            "venues": []
        }


def save_progress(progress):

    with open(
        PROGRESS_FILE,
        "w",
        encoding="utf-8"
    ) as f:

        json.dump(
            progress,
            f,
            ensure_ascii=False,
            indent=2
        )


# ============================================================
# FIND PAGINATION
# ============================================================

def get_page_url(page_number):

    if page_number == 1:
        return VENUES_URL + "/"

    return (
        f"{VENUES_URL}/page/"
        f"{page_number}/"
    )


def get_listing_pages():

    pages = []

    first_url = get_page_url(1)

    soup = get_soup(first_url)

    # --------------------------------------------------------
    # Find pagination links
    # --------------------------------------------------------

    for a in soup.find_all("a", href=True):

        href = normalize_url(
            a.get("href")
        )

        if "/venues/page/" in href:

            pages.append(href)

    pages.append(first_url)

    pages = list(
        dict.fromkeys(pages)
    )

    # --------------------------------------------------------
    # Determine page numbers
    # --------------------------------------------------------

    page_numbers = set()

    for url in pages:

        match = re.search(
            r"/venues/page/(\d+)",
            url
        )

        if match:

            page_numbers.add(
                int(match.group(1))
            )

        else:

            page_numbers.add(1)

    # --------------------------------------------------------
    # If MAX_PAGES is specified
    # --------------------------------------------------------

    if MAX_PAGES:

        page_numbers = {
            x for x in page_numbers
            if x <= MAX_PAGES
        }

    # --------------------------------------------------------
    # Build complete URLs
    # --------------------------------------------------------

    result = []

    for number in sorted(page_numbers):

        result.append(
            get_page_url(number)
        )

    return result


# ============================================================
# EXTRACT VENUE LINKS
# ============================================================

def extract_venue_links(soup):

    links = []

    for a in soup.find_all(
        "a",
        href=True
    ):

        href = a.get("href")

        if not href:
            continue

        full_url = normalize_url(
            href
        )

        parsed = urlparse(full_url)

        # Venue detail URLs are under /venues/
        # but should not be listing pages.
        if (
            parsed.path.startswith("/venues/")
            and "/page/" not in parsed.path
            and parsed.path != "/venues/"
        ):

            links.append(full_url)

    return unique_list(links)


# ============================================================
# GENERIC FIELD EXTRACTION
# ============================================================

def find_text_after_label(
    soup,
    labels
):

    labels = [
        x.lower()
        for x in labels
    ]

    # Look through common elements.
    for element in soup.find_all(
        ["div", "span", "p", "li", "strong", "b"]
    ):

        text = clean_text(
            element.get_text(
                " ",
                strip=True
            )
        )

        lower = text.lower()

        for label in labels:

            if lower.startswith(
                label.lower()
            ):

                remainder = text[
                    len(label):
                ]

                remainder = clean_text(
                    remainder
                )

                if remainder:
                    return remainder

    return ""


# ============================================================
# EXTRACT PRICE
# ============================================================

def extract_price(soup):

    page_text = soup.get_text(
        " ",
        strip=True
    )

    patterns = [

        r"Starting Price\s*[:\-]?\s*(?:PKR|Rs\.?)?\s*([\d,]+)",

        r"Starting From\s*(?:PKR|Rs\.?)?\s*([\d,]+)",

        r"Starting Price.*?([\d,]{4,})",

    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            page_text,
            re.IGNORECASE
        )

        if match:

            value = (
                match.group(1)
                .replace(",", "")
            )

            try:
                return int(value)
            except ValueError:
                pass

    return ""


# ============================================================
# EXTRACT CAPACITY
# ============================================================

def extract_capacity(soup):

    text = clean_text(
        soup.get_text(
            " ",
            strip=True
        )
    )

    patterns = [

        r"Up to\s+([\d,]+)\s*Capacity",

        r"([\d,]+)\s*Capacity",

        r"Capacity\s*[:\-]?\s*([\d,]+)",

    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
            re.IGNORECASE
        )

        if match:

            try:

                return int(
                    match.group(1)
                    .replace(",", "")
                )

            except ValueError:
                pass

    return ""


# ============================================================
# EXTRACT VENUE TYPE + AREA
# ============================================================

def extract_type_and_area(soup):

    venue_type = ""
    area = ""

    # --------------------------------------------------------
    # Search structured text
    # --------------------------------------------------------

    text = clean_text(
        soup.get_text(
            " ",
            strip=True
        )
    )

    # Example:
    # Banquet in Clifton, Karachi

    match = re.search(
        r"\b(Banquet|Ballroom|Rooftop|Marquee|"
        r"Lawn|Indoor|Farmhouse|Restaurant|"
        r"Community Center)\b"
        r".{0,50}?\bin\s+"
        r"([^,]+),\s*Karachi",
        text,
        re.IGNORECASE
    )

    if match:

        venue_type = clean_text(
            match.group(1)
        )

        area = clean_text(
            match.group(2)
        )

    return venue_type, area


# ============================================================
# EXTRACT DESCRIPTION
# ============================================================

def extract_description(soup):

    heading = None

    for element in soup.find_all(
        ["h2", "h3", "h4"]
    ):

        text = clean_text(
            element.get_text(
                " ",
                strip=True
            )
        )

        if "About Venue" in text:

            heading = element
            break

    if heading:

        descriptions = []

        # Collect following elements
        for element in heading.find_all_next():

            if element.name in [
                "h2",
                "h3",
                "h4"
            ]:

                next_heading = clean_text(
                    element.get_text(
                        " ",
                        strip=True
                    )
                )

                if (
                    next_heading
                    and element != heading
                ):
                    break

            if element.name == "p":

                text = clean_text(
                    element.get_text(
                        " ",
                        strip=True
                    )
                )

                if text:
                    descriptions.append(text)

        if descriptions:

            return "\n".join(
                unique_list(
                    descriptions
                )
            )

    return ""


# ============================================================
# EXTRACT AMENITIES / FEATURES
# ============================================================

def extract_features(soup):

    features = []

    heading = None

    for element in soup.find_all(
        ["h2", "h3", "h4"]
    ):

        text = clean_text(
            element.get_text(
                " ",
                strip=True
            )
        )

        if "What this venue offers" in text:

            heading = element
            break

    if heading:

        # Search nearby list items
        parent = heading.parent

        if parent:

            for li in parent.find_all("li"):

                text = clean_text(
                    li.get_text(
                        " ",
                        strip=True
                    )
                )

                if (
                    text
                    and len(text) < 100
                    and text.lower()
                    not in [
                        "image",
                        "show more amenities"
                    ]
                ):

                    features.append(text)

        # Broader fallback
        if not features:

            for element in heading.find_all_next(
                ["li", "span", "div"]
            ):

                text = clean_text(
                    element.get_text(
                        " ",
                        strip=True
                    )
                )

                if (
                    text
                    and len(text) < 80
                    and text.lower()
                    not in [
                        "image",
                        "show more amenities"
                    ]
                ):

                    features.append(text)

                if len(features) >= 30:
                    break

    return unique_list(features)


# ============================================================
# EXTRACT ADDRESS
# ============================================================

def extract_address(soup):

    heading = None

    for element in soup.find_all(
        ["h2", "h3", "h4"]
    ):

        text = clean_text(
            element.get_text(
                " ",
                strip=True
            )
        )

        if "Location & Address" in text:

            heading = element
            break

    if heading:

        # ----------------------------------------------------
        # Look at nearby text
        # ----------------------------------------------------

        for element in heading.find_all_next(
            ["p", "div", "span"]
        ):

            text = clean_text(
                element.get_text(
                    " ",
                    strip=True
                )
            )

            if not text:
                continue

            if (
                "Get Direction" in text
                or "Location & Address" in text
            ):
                continue

            # Avoid huge containers
            if len(text) > 250:
                continue

            # Address usually contains Karachi
            if "karachi" in text.lower():

                return text

    return ""


# ============================================================
# EXTRACT GOOGLE MAPS URL
# ============================================================

def extract_google_maps_url(soup):

    for a in soup.find_all(
        "a",
        href=True
    ):

        href = a.get("href")

        if not href:
            continue

        href_lower = href.lower()

        if (
            "google.com/maps" in href_lower
            or "maps.google.com" in href_lower
            or "goo.gl/maps" in href_lower
        ):

            return href

        text = clean_text(
            a.get_text(
                " ",
                strip=True
            )
        ).lower()

        if (
            "get direction" in text
            and (
                "google" in href_lower
                or "maps" in href_lower
            )
        ):

            return href

    # --------------------------------------------------------
    # Search onclick attributes
    # --------------------------------------------------------

    for element in soup.find_all(
        attrs={
            "onclick": True
        }
    ):

        onclick = element.get(
            "onclick"
        )

        if (
            "google" in onclick.lower()
            and "maps" in onclick.lower()
        ):

            urls = re.findall(
                r'https?://[^\'"\s]+',
                onclick
            )

            if urls:
                return urls[0]

    return ""


# ============================================================
# EXTRACT PHONE / WHATSAPP
# ============================================================

def extract_contact_numbers(soup):

    phones = []
    whatsapp = []

    # --------------------------------------------------------
    # href="tel:"
    # --------------------------------------------------------

    for a in soup.find_all(
        "a",
        href=True
    ):

        href = a.get("href", "")

        text = clean_text(
            a.get_text(
                " ",
                strip=True
            )
        )

        href_lower = href.lower()

        # Telephone
        if href_lower.startswith("tel:"):

            number = href[
                4:
            ].strip()

            phones.append(number)

        # WhatsApp
        if (
            "wa.me" in href_lower
            or "whatsapp.com" in href_lower
        ):

            match = re.search(
                r"(\+?\d{10,15})",
                href
            )

            if match:
                whatsapp.append(
                    match.group(1)
                )

            elif text:
                whatsapp.append(text)

    # --------------------------------------------------------
    # Search entire page for Pakistani numbers
    # --------------------------------------------------------

    page_text = soup.get_text(
        " ",
        strip=True
    )

    phone_patterns = [

        r"\+92[\s\-]?\d{3}[\s\-]?\d{7}",

        r"0092[\s\-]?\d{3}[\s\-]?\d{7}",

        r"03\d{2}[\s\-]?\d{7}",

    ]

    for pattern in phone_patterns:

        matches = re.findall(
            pattern,
            page_text
        )

        phones.extend(matches)

    phones = unique_list(
        phones
    )

    whatsapp = unique_list(
        whatsapp
    )

    return (
        ", ".join(phones),
        ", ".join(whatsapp)
    )


# ============================================================
# EXTRACT IMAGES
# ============================================================

def extract_images(
    soup,
    venue_url
):

    image_urls = []

    # --------------------------------------------------------
    # IMG elements
    # --------------------------------------------------------

    for img in soup.find_all("img"):

        candidates = [

            img.get("src"),

            img.get("data-src"),

            img.get("data-lazy-src"),

            img.get("data-original"),

            img.get("data-image"),

            img.get("data-full"),

        ]

        # srcset
        srcset = img.get(
            "srcset"
        )

        if srcset:

            for part in srcset.split(","):

                url = part.strip().split(
                    " "
                )[0]

                candidates.append(url)

        for candidate in candidates:

            if not candidate:
                continue

            full_url = urljoin(
                venue_url,
                candidate
            )

            if (
                full_url.startswith(
                    "data:"
                )
            ):
                continue

            # Only image-like URLs
            parsed = urlparse(
                full_url
            )

            extension = os.path.splitext(
                parsed.path
            )[1].lower()

            image_extensions = [
                ".jpg",
                ".jpeg",
                ".png",
                ".webp",
                ".gif",
                ".avif"
            ]

            if (
                extension in image_extensions
                or "image" in full_url.lower()
            ):

                image_urls.append(
                    full_url
                )

    # --------------------------------------------------------
    # Background images
    # --------------------------------------------------------

    for element in soup.find_all(
        style=True
    ):

        style = element.get(
            "style",
            ""
        )

        matches = re.findall(
            r'url\([\'"]?([^\'")]+)',
            style
        )

        for image_url in matches:

            image_urls.append(
                urljoin(
                    venue_url,
                    image_url
                )
            )

    # --------------------------------------------------------
    # Gallery links
    # --------------------------------------------------------

    for a in soup.find_all(
        "a",
        href=True
    ):

        href = a.get("href")

        if not href:
            continue

        full_url = urljoin(
            venue_url,
            href
        )

        extension = os.path.splitext(
            urlparse(full_url).path
        )[1].lower()

        if extension in [
            ".jpg",
            ".jpeg",
            ".png",
            ".webp",
            ".gif",
            ".avif"
        ]:

            image_urls.append(
                full_url
            )

    # --------------------------------------------------------
    # Remove duplicates
    # --------------------------------------------------------

    return unique_list(
        image_urls
    )


# ============================================================
# DOWNLOAD IMAGE
# ============================================================

def download_image(
    image_url,
    venue_name,
    image_number
):

    try:

        response = request_get(
            image_url,
            stream=True
        )

        content_type = response.headers.get(
            "Content-Type",
            ""
        ).lower()

        extension = ""

        if "jpeg" in content_type:
            extension = ".jpg"

        elif "png" in content_type:
            extension = ".png"

        elif "webp" in content_type:
            extension = ".webp"

        elif "gif" in content_type:
            extension = ".gif"

        else:

            extension = os.path.splitext(
                urlparse(image_url).path
            )[1].lower()

            if extension not in [
                ".jpg",
                ".jpeg",
                ".png",
                ".webp",
                ".gif",
                ".avif"
            ]:

                extension = ".jpg"

        safe_name = safe_filename(
            venue_name
        )

        # Hash prevents filename collisions
        url_hash = hashlib.md5(
            image_url.encode(
                "utf-8"
            )
        ).hexdigest()[:8]

        filename = (
            f"{safe_name}_"
            f"{image_number}_"
            f"{url_hash}"
            f"{extension}"
        )

        filepath = os.path.join(
            IMAGE_DIR,
            filename
        )

        if os.path.exists(filepath):
            return filepath

        with open(
            filepath,
            "wb"
        ) as f:

            for chunk in response.iter_content(
                chunk_size=64 * 1024
            ):

                if chunk:
                    f.write(chunk)

        return filepath

    except Exception as exc:

        logging.error(
            "Image failed: %s | %s",
            image_url,
            exc
        )

        return ""


# ============================================================
# EXTRACT VENUE
# ============================================================

def scrape_venue(url):

    logging.info(
        "Scraping venue: %s",
        url
    )

    soup = get_soup(url)

    # --------------------------------------------------------
    # Name
    # --------------------------------------------------------

    name = ""

    h1 = soup.find("h1")

    if h1:

        name = clean_text(
            h1.get_text(
                " ",
                strip=True
            )
        )

    # --------------------------------------------------------
    # Type + Area
    # --------------------------------------------------------

    venue_type, area = (
        extract_type_and_area(
            soup
        )
    )

    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    description = (
        extract_description(
            soup
        )
    )

    features = (
        extract_features(
            soup
        )
    )

    address = (
        extract_address(
            soup
        )
    )

    price = (
        extract_price(
            soup
        )
    )

    capacity = (
        extract_capacity(
            soup
        )
    )

    google_maps = (
        extract_google_maps_url(
            soup
        )
    )

    phone, whatsapp = (
        extract_contact_numbers(
            soup
        )
    )

    images = (
        extract_images(
            soup,
            url
        )
    )

    # --------------------------------------------------------
    # Download images
    # --------------------------------------------------------

    downloaded_images = []

    if DOWNLOAD_IMAGES:

        for index, image_url in enumerate(
            images,
            start=1
        ):

            filepath = download_image(
                image_url,
                name,
                index
            )

            if filepath:
                downloaded_images.append(
                    filepath
                )

            time.sleep(
                0.2
            )

    # --------------------------------------------------------
    # Result
    # --------------------------------------------------------

    return {
        "Name": name,
        "Description": description,
        "Address": address,
        "Area": area,
        "Price": price,
        "Capacity": capacity,
        "Venue Type": venue_type,
        "Features": ", ".join(features),
        "Phone": phone,
        "WhatsApp": whatsapp,
        "Google Maps URL": google_maps,
        "Venue URL": url,
        "Image Count": len(images),
        "Image URLs": "\n".join(images),
        "Downloaded Images": "\n".join(
            downloaded_images
        )
    }


# ============================================================
# EXCEL CREATION
# ============================================================

def format_sheet(
    worksheet,
    freeze="A2",
    autofilter=True
):

    worksheet.freeze_panes = freeze

    if autofilter:

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

    # Header
    for cell in worksheet[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # General formatting
    for row in worksheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    # Set useful widths
    widths = {}

    for row in worksheet.iter_rows():

        for cell in row:

            if cell.value is None:
                continue

            value_length = len(
                str(cell.value)
            )

            current = widths.get(
                cell.column,
                0
            )

            widths[cell.column] = min(
                max(
                    current,
                    value_length
                ),
                60
            )

    for column, width in widths.items():

        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = max(
            width,
            12
        )

    worksheet.row_dimensions[1].height = 30


def create_excel(venues):

    wb = Workbook()

    # ========================================================
    # VENUES SHEET
    # ========================================================

    ws = wb.active

    ws.title = "Venues"

    headers = [
        "Name",
        "Description",
        "Address",
        "Area",
        "Price",
        "Capacity",
        "Venue Type",
        "Features",
        "Phone",
        "WhatsApp",
        "Google Maps URL",
        "Venue URL",
        "Image Count",
        "Image URLs",
        "Downloaded Images"
    ]

    ws.append(headers)

    for venue in venues:

        ws.append([
            venue.get(
                "Name",
                ""
            ),

            venue.get(
                "Description",
                ""
            ),

            venue.get(
                "Address",
                ""
            ),

            venue.get(
                "Area",
                ""
            ),

            venue.get(
                "Price",
                ""
            ),

            venue.get(
                "Capacity",
                ""
            ),

            venue.get(
                "Venue Type",
                ""
            ),

            venue.get(
                "Features",
                ""
            ),

            venue.get(
                "Phone",
                ""
            ),

            venue.get(
                "WhatsApp",
                ""
            ),

            venue.get(
                "Google Maps URL",
                ""
            ),

            venue.get(
                "Venue URL",
                ""
            ),

            venue.get(
                "Image Count",
                0
            ),

            venue.get(
                "Image URLs",
                ""
            ),

            venue.get(
                "Downloaded Images",
                ""
            )
        ])

    format_sheet(ws)

    # ========================================================
    # IMAGES SHEET
    # ========================================================

    image_ws = wb.create_sheet(
        "Images"
    )

    image_headers = [
        "Venue Name",
        "Venue URL",
        "Image #",
        "Image URL",
        "Local File"
    ]

    image_ws.append(
        image_headers
    )

    for venue in venues:

        image_urls = venue.get(
            "Image URLs",
            ""
        ).split("\n")

        local_files = venue.get(
            "Downloaded Images",
            ""
        ).split("\n")

        for index, image_url in enumerate(
            image_urls,
            start=1
        ):

            if not image_url:
                continue

            local_file = ""

            if (
                index <= len(
                    local_files
                )
            ):

                local_file = local_files[
                    index - 1
                ]

            image_ws.append([
                venue.get(
                    "Name",
                    ""
                ),

                venue.get(
                    "Venue URL",
                    ""
                ),

                index,

                image_url,

                local_file
            ])

    format_sheet(
        image_ws
    )

    # ========================================================
    # SUMMARY SHEET
    # ========================================================

    summary = wb.create_sheet(
        "Summary",
        0
    )

    summary.append([
        "Metric",
        "Value"
    ])

    total_venues = len(
        venues
    )

    total_images = sum(
        int(
            v.get(
                "Image Count",
                0
            ) or 0
        )
        for v in venues
    )

    summary.append([
        "Total Venues",
        total_venues
    ])

    summary.append([
        "Total Images",
        total_images
    ])

    summary.append([
        "Generated",
        time.strftime(
            "%Y-%m-%d %H:%M:%S"
        )
    ])

    format_sheet(
        summary,
        freeze="A2",
        autofilter=False
    )

    # ========================================================
    # SAVE
    # ========================================================

    wb.save(
        EXCEL_FILE
    )

    logging.info(
        "Excel created: %s",
        EXCEL_FILE
    )


# ============================================================
# MAIN SCRAPER
# ============================================================

def main():

    logging.info(
        "=========================================="
    )

    logging.info(
        "VenueHunt scraper started"
    )

    logging.info(
        "=========================================="
    )

    progress = load_progress()

    scraped_urls = set(
        progress.get(
            "scraped_urls",
            []
        )
    )

    venues = progress.get(
        "venues",
        []
    )

    # --------------------------------------------------------
    # Discover listing pages
    # --------------------------------------------------------

    logging.info(
        "Discovering listing pages..."
    )

    listing_pages = (
        get_listing_pages()
    )

    logging.info(
        "Listing pages found: %s",
        len(listing_pages)
    )

    # --------------------------------------------------------
    # Discover venues
    # --------------------------------------------------------

    venue_urls = []

    for page_url in listing_pages:

        logging.info(
            "Scanning: %s",
            page_url
        )

        try:

            soup = get_soup(
                page_url
            )

            links = (
                extract_venue_links(
                    soup
                )
            )

            for link in links:

                if link not in venue_urls:
                    venue_urls.append(
                        link
                    )

            logging.info(
                "Found %s venue URLs",
                len(links)
            )

            time.sleep(
                REQUEST_DELAY
            )

        except Exception as exc:

            logging.error(
                "Listing page failed: %s | %s",
                page_url,
                exc
            )

    logging.info(
        "Total unique venues discovered: %s",
        len(venue_urls)
    )

    # --------------------------------------------------------
    # Scrape venues
    # --------------------------------------------------------

    for index, venue_url in enumerate(
        venue_urls,
        start=1
    ):

        if venue_url in scraped_urls:

            logging.info(
                "[%s/%s] Skipping already scraped: %s",
                index,
                len(venue_urls),
                venue_url
            )

            continue

        logging.info(
            "[%s/%s] Scraping: %s",
            index,
            len(venue_urls),
            venue_url
        )

        try:

            venue = scrape_venue(
                venue_url
            )

            venues.append(
                venue
            )

            scraped_urls.add(
                venue_url
            )

            # Save after every venue
            progress = {
                "scraped_urls": list(
                    scraped_urls
                ),
                "venues": venues
            }

            save_progress(
                progress
            )

        except Exception as exc:

            logging.error(
                "Venue failed: %s | %s",
                venue_url,
                exc
            )

        time.sleep(
            REQUEST_DELAY
        )

    # --------------------------------------------------------
    # Create Excel
    # --------------------------------------------------------

    create_excel(
        venues
    )

    logging.info(
        "=========================================="
    )

    logging.info(
        "SCRAPING COMPLETED"
    )

    logging.info(
        "Venues: %s",
        len(venues)
    )

    logging.info(
        "Excel: %s",
        EXCEL_FILE
    )

    logging.info(
        "Images: %s",
        IMAGE_DIR
    )

    logging.info(
        "=========================================="
    )


if __name__ == "__main__":
    main()